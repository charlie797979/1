"""뉴스클리핑 핵심 로직: 검색 / 제외 / 중복제거 / 요약 / 엑셀 작성"""
import difflib
import html
import io
import math
import os
import re
from copy import copy
from email.utils import parsedate_to_datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook
from openpyxl.styles import Font

KST = ZoneInfo("Asia/Seoul")
FONT_NAME, FONT_SIZE = "맑은 고딕", 10
SUMMARY_MODEL = os.environ.get("ANTHROPIC_MODEL", "claude-haiku-5-5")
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120 Safari/537.36"}

# 링크 칸 글자 모양 (밑줄·파란색을 빼고 싶으면 LINK_COLOR=None, LINK_UNDERLINE=None)
LINK_COLOR, LINK_UNDERLINE = "0563C1", "single"

STOCK_WORDS = ["주가", "주식", "증시", "코스피", "코스닥", "목표가", "목표주가", "상한가", "하한가",
               "특징주", "종목", "투자의견", "순매수", "순매도", "공매도", "시가총액", "시총", "장중",
               "급등", "급락", "강세", "약세", "증권가", "매수", "매도", "PER", "배당"]
SALES_WORDS = ["매출", "실적", "영업이익"]
GROWTH_WORDS = ["증가", "늘", "성장", "급증", "돌파", "최대", "껑충", "상승", "신장", "↑"]


# ---------------------------------------------------------------- 검색
def keyword_sheets(template_path):
    """엑셀 파일에서 키워드 시트 이름 목록(첫 시트 '월말보고서' 제외)"""
    wb = load_workbook(template_path, read_only=True)
    return wb.sheetnames[1:]


def clean_title(title):
    return html.unescape(re.sub(r"</?b>", "", title or "")).strip()


def parse_naver(items_json, keyword, start_dt, end_dt):
    out = []
    for it in items_json:
        try:
            pub = parsedate_to_datetime(it["pubDate"]).astimezone(KST)
        except Exception:
            continue
        if not (start_dt <= pub <= end_dt):
            continue
        link = it.get("link") or ""
        url = link if "naver.com" in link else (it.get("originallink") or link)
        out.append({"keyword": keyword, "title": clean_title(it.get("title")), "url": url, "pub": pub})
    return out


def search_news(keyword, start_dt, end_dt, client_id, client_secret):
    """네이버 검색 API(뉴스, 최신순). 조회 시작 시각보다 오래된 기사가 나오면 멈춤"""
    headers = {"X-Naver-Client-Id": client_id, "X-Naver-Client-Secret": client_secret}
    out, start = [], 1
    while start <= 1000:
        r = requests.get("https://openapi.naver.com/v1/search/news.json", headers=headers, timeout=20,
                         params={"query": keyword, "display": 100, "start": start, "sort": "date"})
        r.raise_for_status()
        items = r.json().get("items", [])
        if not items:
            break
        out += parse_naver(items, keyword, start_dt, end_dt)
        try:
            oldest = parsedate_to_datetime(items[-1]["pubDate"]).astimezone(KST)
        except Exception:
            break
        if oldest < start_dt or len(items) < 100:
            break
        start += 100
    return out


# ---------------------------------------------------------------- 제외 / 중복
def is_excluded(title, stock=True, sales=True, extra_words=()):
    if stock and any(w in title for w in STOCK_WORDS):
        return True
    if sales and any(w in title for w in SALES_WORDS) and any(w in title for w in GROWTH_WORDS):
        return True
    return any(w and w in title for w in extra_words)


def _norm(title, strip_words=()):
    t = re.sub(r"\[[^\]]*\]", "", title)
    for w in strip_words:  # 검색 키워드는 모든 제목에 들어 있어 유사도를 부풀리므로 비교에서 뺌
        t = t.replace(w, "")
    t = re.sub(r"[^0-9A-Za-z가-힣]", "", t).lower()
    return t or re.sub(r"\s", "", title).lower()


def dedupe(items, threshold=0.5, strip_words=()):
    """제목이 threshold(기본 50%) 이상 비슷하면 같은 기사로 보고 가장 이른 1건만 남김"""
    kept, kept_norm = [], []
    for it in sorted(items, key=lambda x: x["pub"]):
        n = _norm(it["title"], strip_words)
        if any(difflib.SequenceMatcher(None, n, k, autojunk=False).ratio() >= threshold for k in kept_norm):
            continue
        kept.append(it)
        kept_norm.append(n)
    return kept


# ---------------------------------------------------------------- 본문 / 요약
def fetch_body(url):
    try:
        r = requests.get(url, headers=HEADERS, timeout=15)
        if not r.ok:
            return ""
        try:
            import trafilatura
            text = trafilatura.extract(r.content) or ""
        except Exception:
            text = ""
        return text.strip()[:4000]
    except Exception:
        return ""


def fallback_summary(title, body):
    sents = re.split(r"(?<=[.!?다])\s+", body) if body else []
    text = " ".join(sents[:2]).strip()
    return (text[:120] + "…") if len(text) > 120 else (text or title)


def summarize(title, body, api_key=None, model=SUMMARY_MODEL):
    """(요약문, 사용방식) 반환. 사용방식: 'claude' | 'fallback'"""
    if api_key:
        try:
            import anthropic
            client = anthropic.Anthropic(api_key=api_key)
            body_part = body if len(body) >= 200 else "(본문 없음: 제목에 있는 내용만으로 요약하고 없는 내용은 추가하지 말 것)"
            prompt = (
                "다음 뉴스 기사를 한국어 뉴스 문체(~했다, ~이다)로 1~2문장, 공백 포함 80자 이내로 요약하세요. "
                "사실 중심으로 쓰고 수식어·추측은 넣지 마세요. 요약문만 출력하세요.\n\n"
                f"제목: {title}\n본문: {body_part}"
            )
            msg = client.messages.create(model=model, max_tokens=300,
                                         messages=[{"role": "user", "content": prompt}])
            text = "".join(b.text for b in msg.content if b.type == "text").strip()
            if text:
                return text, "claude"
        except Exception:
            pass
    return fallback_summary(title, body), "fallback"


# ---------------------------------------------------------------- 엑셀
def _next_row(ws):
    r = 3
    while r <= ws.max_row and any(ws.cell(r, c).value not in (None, "") for c in range(2, 7)):
        r += 1
    return r


def build_workbook(template_path, rows, work_date):
    """rows: [{keyword, pub, title, url, summary}] -> (xlsx bytes, {키워드: 건수})"""
    wb = load_workbook(template_path)
    plain = Font(name=FONT_NAME, size=FONT_SIZE, bold=False)
    link_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=False, color=LINK_COLOR, underline=LINK_UNDERLINE)
    counts = {}
    for it in rows:
        ws = wb[it["keyword"]]
        r = _next_row(ws)
        for c in range(1, 7):  # 템플릿 서식이 없는 행이면 3행 서식 복사
            cell = ws.cell(r, c)
            if cell.border.left.style is None:
                cell._style = copy(ws.cell(3, c)._style)
        if ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = ws.row_dimensions[3].height
        if ws.cell(r, 1).value is None:
            ws.cell(r, 1).value = r - 2
        ws.cell(r, 2).value = work_date.strftime("%Y.%m.%d")
        ws.cell(r, 3).value = it["pub"].strftime("%Y.%m.%d")
        ws.cell(r, 4).value = it["keyword"]
        ws.cell(r, 5).value = it["summary"]
        link = ws.cell(r, 6)
        link.value = it["title"]
        link.hyperlink = it["url"]
        for c in range(1, 6):
            ws.cell(r, c).font = plain
        link.font = link_font
        lines = math.ceil(len(it["summary"]) / 37)  # 요약 칸 한 줄 ≈ 한글 37자
        if lines > 2:
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 30, 14 * lines)
        counts[it["keyword"]] = counts.get(it["keyword"], 0) + 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), counts
